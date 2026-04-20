
"""
Academic scenario generator for tariff and misinformation inputs.

Approach
--------
This script uses a discrete noisy-signal Bayesian calibration:

1. The user specifies true tariff scenarios with prior probabilities.
2. The user specifies signal states with:
   - perceived tariff (signal anchor)
   - signal probability
   - credibility in [0,1]
   - ambiguity in [0,1]

3. For each signal m and true tariff scenario s, the script computes a likelihood:
      L(m|s) ∝ exp(-lambda_m * |tau_s - hat_tau_m|)
   where
      lambda_m = base_precision * credibility_m * (1 - ambiguity_m)

4. Posterior conditional probabilities are then computed by Bayes' rule:
      Q(s|m) = L(m|s) * prior(s) / Σ_{s'} L(m|s') * prior(s')

This yields the exact model inputs:
- Signals
- TariffScenarios
- ConditionalProb

Run:
    python scenario_generator.py dataset_template.xlsx
"""

from __future__ import annotations
import sys
from pathlib import Path
import math
import pandas as pd
from openpyxl import load_workbook


def soft_likelihood(actual_tau: float, perceived_tau: float, lam: float) -> float:
    return math.exp(-lam * abs(actual_tau - perceived_tau))


def generate_from_workbook(workbook_path: str | Path, base_precision: float = 20.0) -> None:
    workbook_path = Path(workbook_path)

    tariff_priors = pd.read_excel(workbook_path, sheet_name="TariffPriors").dropna(how="all")
    signal_design = pd.read_excel(workbook_path, sheet_name="SignalDesign").dropna(how="all")

    tariff_priors.columns = [str(c).strip() for c in tariff_priors.columns]
    signal_design.columns = [str(c).strip() for c in signal_design.columns]

    tariff_priors["scenario_id"] = tariff_priors["scenario_id"].astype(str)
    signal_design["signal_id"] = signal_design["signal_id"].astype(str)

    # Normalize priors and signal probabilities defensively
    tariff_priors["prior_prob"] = tariff_priors["prior_prob"] / tariff_priors["prior_prob"].sum()
    signal_design["signal_prob"] = signal_design["signal_prob"] / signal_design["signal_prob"].sum()

    # Create Signals and TariffScenarios output tables
    signals_out = signal_design[["signal_id", "signal_prob", "perceived_tariff"]].copy()
    tariff_out = tariff_priors[["scenario_id", "actual_tariff"]].copy()

    cond_rows = []
    for _, mrow in signal_design.iterrows():
        m = str(mrow["signal_id"])
        perceived_tau = float(mrow["perceived_tariff"])
        credibility = float(mrow["credibility"])
        ambiguity = float(mrow["ambiguity"])

        lam = max(1e-6, base_precision * credibility * (1.0 - ambiguity))

        numerators = []
        for _, srow in tariff_priors.iterrows():
            s = str(srow["scenario_id"])
            actual_tau = float(srow["actual_tariff"])
            prior_prob = float(srow["prior_prob"])
            likelihood = soft_likelihood(actual_tau, perceived_tau, lam)
            numerators.append((s, likelihood * prior_prob))

        denom = sum(v for _, v in numerators)
        for s, num in numerators:
            cond_rows.append({
                "signal_id": m,
                "scenario_id": s,
                "cond_prob": num / denom if denom > 0 else 1.0 / len(numerators)
            })

    cond_out = pd.DataFrame(cond_rows)

    # Write back to workbook, replacing generated sheets
    wb = load_workbook(workbook_path)
    for sheet_name in ["Signals", "TariffScenarios", "ConditionalProb"]:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            wb.remove(ws)

    ws_sig = wb.create_sheet("Signals")
    ws_sig.append(list(signals_out.columns))
    for row in signals_out.itertuples(index=False):
        ws_sig.append(list(row))

    ws_tar = wb.create_sheet("TariffScenarios")
    ws_tar.append(list(tariff_out.columns))
    for row in tariff_out.itertuples(index=False):
        ws_tar.append(list(row))

    ws_cond = wb.create_sheet("ConditionalProb")
    ws_cond.append(list(cond_out.columns))
    for row in cond_out.itertuples(index=False):
        ws_cond.append(list(row))

    wb.save(workbook_path)

    print(f"Generated scenario sheets in: {workbook_path}")


if __name__ == "__main__":
    # if len(sys.argv) < 2:
    #     raise SystemExit("Usage: python scenario_generator.py <dataset_template.xlsx> [base_precision]")
    # workbook = sys.argv[1]
    # base_precision = float(sys.argv[2]) if len(sys.argv) >= 3 else 20.0
    
    if len(sys.argv) < 2:
        workbook = "dataset_template.xlsx"
        base_precision = 20.0
    else:
        workbook = sys.argv[1]
        base_precision = float(sys.argv[2]) if len(sys.argv) >= 3 else 20.0
    generate_from_workbook(workbook, base_precision=base_precision)
